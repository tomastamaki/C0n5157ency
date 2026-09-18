"""
Parsea el Excel del programa "Min-Max Phase 2: Peak Physique (4x/semana)" y genera
src/data/program.json, que la app React consume directamente.

Uso:
    python scripts/parse_program.py [ruta-al-excel]

Si no se pasa ruta, busca "Min-Max_Phase_2_-_4x.xlsx" en la raíz del repo.

Volvé a correr este script cada vez que cambies de fase/programa y reemplaces el Excel
de origen (o pases la ruta de un Excel distinto con el mismo layout de columnas).

Requiere: pip install openpyxl
"""

import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl

SHEET_NAME = "4x Per Week"
DAY_NAMES = {"Upper", "Lower", "Push", "Pull"}

# Columnas (1-indexado, A=1)
COL = {
    "day": 2,          # B
    "exercise": 3,      # C
    "intensity": 4,      # D
    "warmup": 5,      # E
    "working": 6,      # F
    "reps": 7,      # G
    "rir1": 14,     # N
    "rir2": 15,     # O
    "rir3": 16,     # P
    "rest": 17,     # Q
    "sub1": 18,     # R
    "sub2": 19,     # S
    "notes": 20,     # T
}

BLOCK_RE = re.compile(r"^Block\s+\d+$", re.IGNORECASE)
WEEK_RE = re.compile(r"^Week\s+(\d+)$", re.IGNORECASE)
SUPERSET_RE = re.compile(r"^S(\d+):\s*(.+)$")


def to_range_string(cell):
    """Recupera el string original (ej. '4-6') aunque Excel lo haya convertido en fecha."""
    v = cell.value
    if v is None:
        return None
    if isinstance(v, datetime.datetime):
        return f"{v.month}-{v.day}"
    if isinstance(v, (int, float)):
        return str(int(v)) if float(v).is_integer() else str(v)
    return str(v).strip()


def to_int(cell):
    v = cell.value
    if v is None:
        return None
    if isinstance(v, (int, float)):
        return int(v)
    try:
        return int(str(v).strip())
    except ValueError:
        return None


def to_rir_str(cell):
    v = cell.value
    if v is None:
        return None
    if isinstance(v, str) and v.strip() == "-":
        return None
    if isinstance(v, (int, float)):
        return str(int(v))
    return str(v).strip()


def to_text_or_none(cell, dash_is_none=True):
    v = cell.value
    if v is None:
        return None
    s = str(v).strip()
    if not s or (dash_is_none and s == "-"):
        return None
    return s


def to_embed_url(url):
    if not url:
        return None
    m = re.search(r"youtu\.be/([\w-]+)", url)
    if not m:
        m = re.search(r"[?&]v=([\w-]+)", url)
    if not m:
        return url
    return f"https://www.youtube.com/embed/{m.group(1)}"


def parse_exercise_row(ws, row):
    raw_name = to_text_or_none(ws.cell(row=row, column=COL["exercise"]), dash_is_none=False)
    m = SUPERSET_RE.match(raw_name)
    if m:
        superset_group = f"S{m.group(1)}"
        name = m.group(2).strip()
    else:
        superset_group = None
        name = raw_name

    exercise_cell = ws.cell(row=row, column=COL["exercise"])
    video_url = to_embed_url(exercise_cell.hyperlink.target if exercise_cell.hyperlink else None)

    warmup = to_range_string(ws.cell(row=row, column=COL["warmup"]))
    working_count = to_int(ws.cell(row=row, column=COL["working"]))
    reps = to_range_string(ws.cell(row=row, column=COL["reps"]))

    target_rir = []
    for key in ("rir1", "rir2", "rir3"):
        rir = to_rir_str(ws.cell(row=row, column=COL[key]))
        if rir is not None:
            target_rir.append(rir)

    substitutions = []
    for key in ("sub1", "sub2"):
        sub_cell = ws.cell(row=row, column=COL[key])
        sub_name = to_text_or_none(sub_cell)
        if sub_name:
            substitutions.append(
                {
                    "name": sub_name,
                    "videoUrl": to_embed_url(sub_cell.hyperlink.target if sub_cell.hyperlink else None),
                }
            )

    return {
        "name": name,
        "supersetGroup": superset_group,
        "videoUrl": video_url,
        "intensityTechnique": to_text_or_none(ws.cell(row=row, column=COL["intensity"])),
        "notes": to_text_or_none(ws.cell(row=row, column=COL["notes"])),
        "substitutions": substitutions,
        "rest": to_text_or_none(ws.cell(row=row, column=COL["rest"])),
        "warmup": warmup,
        "workingCount": working_count,
        "reps": reps,
        "targetRIR": target_rir,
    }


def is_zero_or_none(warmup_str):
    return warmup_str is None or warmup_str.strip() == "0"


def append_exercise_row(day, rec):
    groups = day["exerciseGroups"]
    merge_key = (rec["name"], rec["supersetGroup"])
    if groups and groups[-1]["_mergeKey"] == merge_key:
        group = groups[-1]
    else:
        group = {
            "exercise": rec["name"],
            "videoUrl": rec["videoUrl"],
            "supersetGroup": rec["supersetGroup"],
            "intensityTechnique": rec["intensityTechnique"],
            "notes": rec["notes"],
            "substitutions": rec["substitutions"],
            "rest": rec["rest"],
            "sets": [],
            "_mergeKey": merge_key,
        }
        groups.append(group)

    if not group["intensityTechnique"] and rec["intensityTechnique"]:
        group["intensityTechnique"] = rec["intensityTechnique"]
    if not group["notes"] and rec["notes"]:
        group["notes"] = rec["notes"]

    if not is_zero_or_none(rec["warmup"]):
        group["sets"].append({"type": "warmup", "count": rec["warmup"]})

    if len(rec["targetRIR"]) != (rec["workingCount"] or 0):
        print(
            f"  [warn] fila con {rec['workingCount']} working sets pero "
            f"{len(rec['targetRIR'])} valores de RIR: '{rec['name']}'",
            file=sys.stderr,
        )

    group["sets"].append(
        {"type": "working", "reps": rec["reps"], "targetRIR": rec["targetRIR"]}
    )


def parse_workbook(path):
    wb = openpyxl.load_workbook(path, data_only=False)
    ws = wb[SHEET_NAME]

    blocks = []
    current_block = None
    current_week = None
    current_day = None
    pending_label = None

    for row in range(1, ws.max_row + 1):
        b_cell = ws.cell(row=row, column=COL["day"])
        b = b_cell.value
        c = ws.cell(row=row, column=COL["exercise"]).value

        if isinstance(b, str) and b.strip():
            bs = b.strip()

            if BLOCK_RE.match(bs):
                current_block = {"name": bs, "weeks": []}
                blocks.append(current_block)
                current_week = None
                current_day = None
                pending_label = None
                continue

            week_match = WEEK_RE.match(bs)
            if week_match:
                current_week = {
                    "label": pending_label or bs,
                    "weekNumber": int(week_match.group(1)),
                    "days": [],
                }
                current_block["weeks"].append(current_week)
                current_day = None
                pending_label = None
                continue

            if bs in DAY_NAMES:
                current_day = {"name": bs, "exerciseGroups": []}
                current_week["days"].append(current_day)
                # esta misma fila también trae el primer ejercicio del día (procesado abajo)
            elif "rest day" in bs.lower():
                pending_label = None
                continue
            else:
                # etiqueta especial de semana, ej. "Intro Week" / "Deload Week"
                pending_label = bs
                continue

        if current_day is not None and isinstance(c, str) and c.strip() and c.strip() != "Exercise":
            rec = parse_exercise_row(ws, row)
            append_exercise_row(current_day, rec)

    # limpiar claves internas de merge
    for block in blocks:
        for week in block["weeks"]:
            for day in week["days"]:
                for group in day["exerciseGroups"]:
                    group.pop("_mergeKey", None)

    return {"blocks": blocks}


def main():
    repo_root = Path(__file__).resolve().parent.parent
    if len(sys.argv) > 1:
        xlsx_path = Path(sys.argv[1])
    else:
        xlsx_path = repo_root / "Min-Max_Phase_2_-_4x.xlsx"

    if not xlsx_path.exists():
        print(f"No se encontró el Excel en: {xlsx_path}", file=sys.stderr)
        sys.exit(1)

    print(f"Parseando {xlsx_path} ...")
    program = parse_workbook(xlsx_path)

    n_weeks = sum(len(b["weeks"]) for b in program["blocks"])
    n_days = sum(len(w["days"]) for b in program["blocks"] for w in b["weeks"])
    n_exercises = sum(
        len(d["exerciseGroups"])
        for b in program["blocks"]
        for w in b["weeks"]
        for d in w["days"]
    )
    print(
        f"  {len(program['blocks'])} bloques, {n_weeks} semanas, "
        f"{n_days} días, {n_exercises} grupos de ejercicio"
    )

    out_path = repo_root / "src" / "data" / "program.json"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(program, indent=2, ensure_ascii=False), encoding="utf-8")
    print(f"Escrito {out_path}")


if __name__ == "__main__":
    main()
